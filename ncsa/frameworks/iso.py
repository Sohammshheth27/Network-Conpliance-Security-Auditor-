"""ISO/IEC 27001:2022 via the official NIST OLIR crosswalk.

Plan 6.2's ISO shortcut:

    our control -> NIST 800-53 ID -> [official NIST crosswalk] -> ISO clause

We never invent an ISO mapping, and we never need to buy the ISO standard to
produce a complete ISO answer for every vendor.

On copyright: the crosswalk ships clause NUMBERS with an empty description
column -- NIST omitted ISO's text deliberately. Verified on the real file:
0 of 177 sampled rows carry any ISO description. So there is nothing of ISO's
to discard here; the mapping itself is complete and free to use. Annex A clause
names below are the short public labels used in the standard's own contents
listing (plan 6.2 permits "clause number + short title").

Source: csrc.nist.gov OLIR, sp800-53r5-to-iso-27001-mapping-2022.
"""
from pathlib import Path

from .models import Automatability, Catalog, CatalogEntry, Framework, License

# Annex A short titles -- clause number + short name only, per plan 6.2.
# Used so a report can say "A.8.20 Network security" rather than a bare number.
ANNEX_A_TITLES: dict[str, str] = {
    "A.5.1": "Policies for information security",
    "A.5.2": "Information security roles and responsibilities",
    "A.5.3": "Segregation of duties",
    "A.5.7": "Threat intelligence",
    "A.5.9": "Inventory of information and other associated assets",
    "A.5.10": "Acceptable use of information and other associated assets",
    "A.5.12": "Classification of information",
    "A.5.14": "Information transfer",
    "A.5.15": "Access control",
    "A.5.16": "Identity management",
    "A.5.17": "Authentication information",
    "A.5.18": "Access rights",
    "A.5.23": "Information security for use of cloud services",
    "A.5.28": "Collection of evidence",
    "A.5.31": "Legal, statutory, regulatory and contractual requirements",
    "A.5.33": "Protection of records",
    "A.5.35": "Independent review of information security",
    "A.5.36": "Compliance with policies, rules and standards",
    "A.5.37": "Documented operating procedures",
    "A.6.3": "Information security awareness, education and training",
    "A.7.1": "Physical security perimeters",
    "A.7.4": "Physical security monitoring",
    "A.8.1": "User endpoint devices",
    "A.8.2": "Privileged access rights",
    "A.8.3": "Information access restriction",
    "A.8.4": "Access to source code",
    "A.8.5": "Secure authentication",
    "A.8.6": "Capacity management",
    "A.8.7": "Protection against malware",
    "A.8.8": "Management of technical vulnerabilities",
    "A.8.9": "Configuration management",
    "A.8.10": "Information deletion",
    "A.8.11": "Data masking",
    "A.8.12": "Data leakage prevention",
    "A.8.13": "Information backup",
    "A.8.15": "Logging",
    "A.8.16": "Monitoring activities",
    "A.8.17": "Clock synchronization",
    "A.8.18": "Use of privileged utility programs",
    "A.8.19": "Installation of software on operational systems",
    "A.8.20": "Network security",
    "A.8.21": "Security of network services",
    "A.8.22": "Segregation of networks",
    "A.8.23": "Web filtering",
    "A.8.24": "Use of cryptography",
    "A.8.25": "Secure development life cycle",
    "A.8.26": "Application security requirements",
    "A.8.28": "Secure coding",
    "A.8.29": "Security testing in development and acceptance",
    "A.8.31": "Separation of development, test and production environments",
    "A.8.32": "Change management",
    "A.8.34": "Protection of information systems during audit testing",
}

# Clauses 4-10 are the management-system requirements: always procedural.
_MGMT_CLAUSE_PREFIXES = tuple(str(n) for n in range(4, 11))


def load(xlsx_path: str | Path) -> tuple[Catalog, dict[str, list[str]]]:
    """Return (ISO catalogue, NIST-ID -> [ISO clause] mapping).

    The mapping is the operational half: it is what turns a NIST label on one
    of our controls into a complete ISO answer.
    """
    import openpyxl

    p = Path(xlsx_path)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)

    mapping: dict[str, list[str]] = {}
    clauses: set[str] = set()

    for sheet in wb.sheetnames:
        if sheet.lower().startswith("definition"):
            continue
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 4:
                continue
            nist_id = str(row[0]).strip().upper() if row[0] else ""
            clause = str(row[3]).strip() if row[3] else ""
            if not nist_id or not clause or clause.lower() == "none":
                continue
            # OSCAL writes AC-1; the crosswalk writes AC-01. Normalise to OSCAL.
            nist_id = _normalise_nist_id(nist_id)
            mapping.setdefault(nist_id, [])
            if clause not in mapping[nist_id]:
                mapping[nist_id].append(clause)
            clauses.add(clause)

    entries = [
        CatalogEntry(
            framework=Framework.ISO_27001,
            id=c,
            title=ANNEX_A_TITLES.get(c),        # short public label, or None
            license=License.IDENTIFIER_ONLY,     # never paste ISO prose
            automatable=(
                Automatability.PROCEDURAL
                if c.startswith(_MGMT_CLAUSE_PREFIXES)
                else Automatability.UNKNOWN
            ),
            source_document="ISO/IEC 27001:2022 (via NIST OLIR crosswalk)",
            source_file=p.name,
            extra={"annex_a": c.startswith("A."), "redistributable": False},
        )
        for c in sorted(clauses, key=_clause_sort_key)
    ]

    return (
        Catalog(framework=Framework.ISO_27001, entries=entries, sources=[str(p)]),
        mapping,
    )


def _normalise_nist_id(cid: str) -> str:
    """Canonicalise a NIST control id to OSCAL's dot form.

    The two sources disagree on notation, and the mismatch is not cosmetic:
    OSCAL writes enhancements as ``AC-2.1`` while the OLIR crosswalk writes
    ``AC-2(12)``. 872 of the 1196 OSCAL entries are enhancements, so leaving
    the formats unreconciled meant none of them could ever match a crosswalk
    row -- ISO coverage silently capped at the base controls.

        AC-01      -> AC-2 style: AC-1
        AC-02(01)  -> AC-2.1
        AC-2.1     -> AC-2.1  (already canonical)
    """
    import re

    m = re.match(r"^([A-Z]{2})-0*(\d+)(?:\s*[.(]0*(\d+)\)?)?$", cid.strip())
    if not m:
        return cid
    fam, num, enh = m.groups()
    return f"{fam}-{int(num)}" + (f".{int(enh)}" if enh else "")


def _clause_sort_key(c: str):
    import re

    annex = c.startswith("A.")
    nums = tuple(int(x) for x in re.findall(r"\d+", c))
    return (annex, nums)
